#!/usr/bin/env python3
"""Generate data/unit-manifest.js from the locked assignment workbook."""

from collections import defaultdict
from hashlib import sha256
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "sources" / "考研阅读1500核心词_60组固定归属_V1.2_最终冻结候选版.xlsx"
TARGET = ROOT / "data" / "unit-manifest.js"


def main():
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    rows = workbook["60组固定归属"].iter_rows(values_only=True)
    headers = next(rows)
    index = {name: position for position, name in enumerate(headers)}
    units = defaultdict(list)
    stages = {}
    for row in rows:
        if not any(value is not None for value in row):
            continue
        unit_id = str(row[index["Unit"]]).strip()
        units[unit_id].append(str(row[index["目标词"]]).strip().lower())
        stages[unit_id] = str(row[index["阶段"]]).strip()

    overview_rows = workbook["Unit总览"].iter_rows(values_only=True)
    overview_headers = next(overview_rows)
    overview_index = {name: position for position, name in enumerate(overview_headers)}
    overview = {
        str(row[overview_index["Unit"]]).strip(): {
            "themeAnchor": str(row[overview_index["主题锚点"]]).strip(),
            "compatibilityRisk": str(row[overview_index["状态"]]).strip(),
        }
        for row in overview_rows
        if any(value is not None for value in row)
    }

    entries = {
        unit_id: {
            "id": unit_id,
            "sequence": int(unit_id[1:]),
            "status": "locked",
            "stage": stages[unit_id],
            "themeAnchor": overview.get(unit_id, {}).get("themeAnchor", "U01冻结"),
            "compatibilityRisk": overview.get(unit_id, {}).get("compatibilityRisk", "冻结"),
            "officialTargetWords": units[unit_id],
        }
        for unit_id in sorted(units)
    }
    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    entries_json = json.dumps(entries, ensure_ascii=False, indent=2)
    content = f'''// Generated from the locked Excel source by scripts/build_unit_manifest.py. Do not hand-edit assignments.
(function () {{
  const UNIT_COUNT = 60;
  const WORDS_PER_UNIT = 25;
  const TOTAL_WORDS = 1500;
  const normalizeWord = word => String(word || "").trim().toLowerCase();
  const source = Object.freeze({{
    version: "V1.2",
    workbook: "考研阅读1500核心词_60组固定归属_V1.2_最终冻结候选版.xlsx",
    sheet: "60组固定归属",
    sha256: "{source_hash}"
  }});
  const entries = Object.freeze({entries_json});

  function duplicates(words) {{
    const seen = new Set();
    const repeated = new Set();
    words.forEach(word => seen.has(word) ? repeated.add(word) : seen.add(word));
    return [...repeated];
  }}

  function inspect() {{
    const errors = [];
    const expectedIds = Array.from({{ length: UNIT_COUNT }}, (_, index) => `U${{String(index + 1).padStart(2, "0")}}`);
    const ids = Object.keys(entries);
    const missingIds = expectedIds.filter(id => !entries[id]);
    const extraIds = ids.filter(id => !expectedIds.includes(id));
    if (ids.length !== UNIT_COUNT) errors.push(`正式 Manifest Unit 数量应为 ${{UNIT_COUNT}}，当前为 ${{ids.length}}。`);
    if (missingIds.length) errors.push(`Manifest 缺少 Unit：${{missingIds.join(", ")}}。`);
    if (extraIds.length) errors.push(`Manifest 存在非法 Unit：${{extraIds.join(", ")}}。`);

    const owners = new Map();
    ids.forEach(id => {{
      const entry = entries[id];
      const words = Array.isArray(entry.officialTargetWords) ? entry.officialTargetWords.map(normalizeWord) : [];
      const emptyIndexes = words.flatMap((word, index) => word ? [] : [index + 1]);
      const repeated = duplicates(words.filter(Boolean));
      if (entry.id !== id) errors.push(`${{id}} 的 entry.id 不一致。`);
      if (entry.sequence !== Number(id.slice(1))) errors.push(`${{id}} 的 sequence 不正确。`);
      if (entry.status !== "locked") errors.push(`${{id}} 尚未锁定。`);
      if (words.length !== WORDS_PER_UNIT) errors.push(`${{id}} 正式目标词应为 ${{WORDS_PER_UNIT}} 个，当前为 ${{words.length}} 个。`);
      if (emptyIndexes.length) errors.push(`${{id}} 存在空词项，位置：${{emptyIndexes.join(", ")}}。`);
      if (repeated.length) errors.push(`${{id}} Unit 内重复词：${{repeated.join(", ")}}。`);
      words.filter(Boolean).forEach(word => {{
        if (!owners.has(word)) owners.set(word, []);
        owners.get(word).push(id);
      }});
    }});
    const crossUnitDuplicates = [...owners].filter(([, unitIds]) => unitIds.length > 1);
    if (owners.size !== TOTAL_WORDS) errors.push(`全局正式目标词应为 ${{TOTAL_WORDS}} 个唯一词，当前为 ${{owners.size}} 个。`);
    if (crossUnitDuplicates.length) errors.push(`Unit 间重复词：${{crossUnitDuplicates.map(([word, ids]) => `${{word}} (${{ids.join("/")}})`).join(", ")}}。`);
    return {{ valid: errors.length === 0, errors, lockedUnitCount: ids.length, lockedWordCount: owners.size, pendingUnits: [] }};
  }}

  function validateOfficialUnit(unit) {{
    if (unit.status !== "official") return;
    const entry = entries[unit.id];
    if (!entry) throw new Error(`正式 Unit ${{unit.id}} 不存在于 Unit Manifest。`);
    const manifestWords = entry.officialTargetWords.map(normalizeWord);
    const unitWords = unit.vocabulary.map(item => normalizeWord(item?.word));
    const manifestSet = new Set(manifestWords);
    const unitSet = new Set(unitWords);
    const missingWords = manifestWords.filter(word => !unitSet.has(word));
    const extraWords = [...unitSet].filter(word => !manifestSet.has(word));
    const repeated = duplicates(unitWords.filter(Boolean));
    const problems = [];
    if (unitWords.length !== WORDS_PER_UNIT) problems.push(`数量应为 ${{WORDS_PER_UNIT}}，当前为 ${{unitWords.length}}`);
    if (missingWords.length) problems.push(`缺少：${{missingWords.join(", ")}}`);
    if (extraWords.length) problems.push(`多出/换入：${{extraWords.join(", ")}}`);
    if (repeated.length) problems.push(`重复：${{repeated.join(", ")}}`);
    if (problems.length) throw new Error(`${{unit.id}} vocabulary 与正式 Manifest 不一致：${{problems.join("；")}}。`);
  }}

  window.UnitManifest = Object.freeze({{
    unitCount: UNIT_COUNT,
    wordsPerUnit: WORDS_PER_UNIT,
    totalWords: TOTAL_WORDS,
    source,
    entries,
    inspect,
    assertComplete() {{
      const result = inspect();
      if (result.errors.length) throw new Error(`正式 Unit Manifest 未通过校验：\\n${{result.errors.join("\\n")}}`);
      return result;
    }},
    validateOfficialUnit
  }});
}})();
'''
    TARGET.write_text(content, encoding="utf-8", newline="\n")
    print(f"Generated {TARGET} from {SOURCE.name} ({len(entries)} Units).")


if __name__ == "__main__":
    main()
