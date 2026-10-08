#!/usr/bin/env python3
"""Read-only validator for the locked 60-Unit assignment workbook.

The script never repairs or rewrites source data. Any mismatch exits non-zero.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from pathlib import Path
import sys

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ASSIGNMENT = ROOT / "data" / "sources" / "考研阅读1500核心词_60组固定归属_V1.2_最终冻结候选版.xlsx"
DEFAULT_MASTER = ROOT / "data" / "sources" / "2005-2025_考研阅读1500核心词母库_V1.0_正式锁定版.xlsx"

EXPECTED_SHEETS = ("60组固定归属", "Unit总览", "V1.2局部调整记录", "冻结校验")
EXPECTED_UNITS = [f"U{i:02d}" for i in range(1, 61)]
FROZEN_U01_WORDS = {
    "criterion", "produce", "afford", "bargain", "expense", "picture",
    "harm", "manage", "meet", "merit", "oblige", "ownership", "scrap",
    "spark", "supply", "alternative", "sustain", "wear", "worn",
    "solution", "decrease", "efficiency", "lifetime", "minimum", "reasonable",
}
FROZEN_U02_WORDS = {
    "exist", "standard", "board", "behind", "outside", "join", "drive",
    "majority", "basic", "line", "editor", "panel", "extra", "particular",
    "professional", "check", "medical", "welcome", "cell", "officer", "bag",
    "beginning", "win", "stick", "additional",
}
FROZEN_U56_WORDS = {
    "expensive", "spend", "excessive", "undergraduate", "fee", "production",
    "path", "tempt", "sight", "dedicate", "bed", "cash", "decorate", "employ",
    "ground", "luxury", "meat", "mobile", "nasty", "notice", "pace", "painting",
    "plenty", "possession", "predecessor",
}


def normalized(value: object) -> str:
    return str(value or "").strip().lower()


def records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]


def duplicates(values):
    return sorted(value for value, count in Counter(values).items() if value and count > 1)


def validate(assignment_path: Path, master_path: Path):
    errors = []
    assignment = openpyxl.load_workbook(assignment_path, read_only=True, data_only=True)
    missing_sheets = [name for name in EXPECTED_SHEETS if name not in assignment.sheetnames]
    if missing_sheets:
        errors.append(f"缺少指定 Sheet：{', '.join(missing_sheets)}")
        return errors, {}

    rows = records(assignment["60组固定归属"])
    overview = records(assignment["Unit总览"])
    adjustments = records(assignment["V1.2局部调整记录"])
    freeze_checks = records(assignment["冻结校验"])
    by_unit = defaultdict(list)
    for index, row in enumerate(rows, start=2):
        unit_id = str(row.get("Unit") or "").strip()
        word = normalized(row.get("目标词"))
        if not unit_id:
            errors.append(f"60组固定归属!A{index}：Unit 为空")
        if not word:
            errors.append(f"60组固定归属!D{index}：目标词为空")
        by_unit[unit_id].append((word, row, index))

    actual_units = sorted(unit_id for unit_id in by_unit if unit_id)
    missing_units = [unit_id for unit_id in EXPECTED_UNITS if unit_id not in by_unit]
    extra_units = [unit_id for unit_id in actual_units if unit_id not in EXPECTED_UNITS]
    if len(actual_units) != 60:
        errors.append(f"Unit 数量应为60，当前为{len(actual_units)}")
    if missing_units:
        errors.append(f"Unit 缺号：{', '.join(missing_units)}")
    if extra_units:
        errors.append(f"非法 Unit 编号：{', '.join(extra_units)}")

    for unit_id in EXPECTED_UNITS:
        items = by_unit.get(unit_id, [])
        words = [item[0] for item in items]
        if len(words) != 25:
            errors.append(f"{unit_id} 应有25词，当前为{len(words)}")
        repeated = duplicates(words)
        if repeated:
            errors.append(f"{unit_id} 组内重复：{', '.join(repeated)}")
        positions = [item[1].get("组内序号") for item in items]
        if positions != list(range(1, 26)):
            errors.append(f"{unit_id} 组内序号不是连续的1–25")

    all_words = [item[0] for items in by_unit.values() for item in items]
    if len(all_words) != 1500:
        errors.append(f"总目标词应为1500，当前为{len(all_words)}")
    repeated_globally = duplicates(all_words)
    if repeated_globally:
        errors.append(f"Unit 间重复词：{', '.join(repeated_globally)}")

    u01_words = {item[0] for item in by_unit.get("U01", [])}
    if u01_words != FROZEN_U01_WORDS:
        errors.append(
            "U01 与冻结版不一致；缺少：{}；多出：{}".format(
                ", ".join(sorted(FROZEN_U01_WORDS - u01_words)) or "无",
                ", ".join(sorted(u01_words - FROZEN_U01_WORDS)) or "无",
            )
        )

    overview_units = [str(row.get("Unit") or "").strip() for row in overview]
    if overview_units != EXPECTED_UNITS:
        errors.append("Unit总览中的 Unit 编号不是连续 U01–U60")
    for row in overview:
        if row.get("词数") != 25:
            errors.append(f"Unit总览：{row.get('Unit')} 词数不是25")

    u02_words = {item[0] for item in by_unit.get("U02", [])}
    if u02_words != FROZEN_U02_WORDS:
        errors.append("U02 未保持上一轮 PASS 版本")
    u56_words = {item[0] for item in by_unit.get("U56", [])}
    if u56_words != FROZEN_U56_WORDS:
        errors.append("U56 未保持上一轮版本")

    master = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    if "1500核心词母库" not in master.sheetnames:
        errors.append("正式母库缺少 Sheet：1500核心词母库")
    else:
        master_rows = records(master["1500核心词母库"])
        master_by_sequence = {row.get("最终序号"): normalized(row.get("统计词元")) for row in master_rows}
        master_words = [normalized(row.get("统计词元")) for row in master_rows]
        if len(master_words) != 1500 or len(set(master_words)) != 1500:
            errors.append("正式母库本身未形成1500个唯一统计词元")
        for word, row, excel_row in [item for items in by_unit.values() for item in items]:
            sequence = row.get("母库最终序号")
            expected_word = master_by_sequence.get(sequence)
            if expected_word != word:
                errors.append(
                    f"60组固定归属!D{excel_row}：{word or '<空>'} 与母库序号 {sequence} 对应词 {expected_word or '<不存在>'} 不一致"
                )
            if not normalized(row.get("5500官方词形")):
                errors.append(f"60组固定归属!E{excel_row}：5500官方词形为空")

    result = {
        "units": len(actual_units),
        "words": len(all_words),
        "perUnitCounts": sorted({len(by_unit.get(unit_id, [])) for unit_id in EXPECTED_UNITS}),
        "duplicates": repeated_globally,
        "u01MatchesFrozen": u01_words == FROZEN_U01_WORDS,
        "overviewRows": len(overview),
        "adjustmentRows": len(adjustments),
        "freezeCheckRows": len(freeze_checks),
        "u02Unchanged": u02_words == FROZEN_U02_WORDS,
        "u56Unchanged": u56_words == FROZEN_U56_WORDS,
        "masterChecked": True,
    }
    return errors, result


def main():
    parser = argparse.ArgumentParser(description="Validate the locked 60×25 Unit assignment without modifying it.")
    parser.add_argument("--assignment", type=Path, default=DEFAULT_ASSIGNMENT)
    parser.add_argument("--master", type=Path, default=DEFAULT_MASTER)
    args = parser.parse_args()
    for path in (args.assignment, args.master):
        if not path.is_file():
            print(f"ERROR: 文件不存在：{path}", file=sys.stderr)
            return 2
    errors, result = validate(args.assignment, args.master)
    if errors:
        print("UNIT ASSIGNMENT VALIDATION: FAIL", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    print("UNIT ASSIGNMENT VALIDATION: PASS")
    print(f"- Units: {result['units']} / 60")
    print(f"- Words per Unit: {result['perUnitCounts']} (expected [25])")
    print(f"- Total official target words: {result['words']} / 1500")
    print(f"- Global duplicates: {len(result['duplicates'])}")
    print(f"- U01 frozen assignment: {'PASS' if result['u01MatchesFrozen'] else 'FAIL'}")
    print(f"- Official 1500 master cross-check: {'PASS' if result['masterChecked'] else 'FAIL'}")
    print(f"- U02 previous PASS assignment unchanged: {'PASS' if result['u02Unchanged'] else 'FAIL'}")
    print(f"- U56 previous assignment unchanged: {'PASS' if result['u56Unchanged'] else 'FAIL'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
